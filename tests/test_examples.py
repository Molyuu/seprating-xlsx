import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from splitter.engine import Options, analyze, default_output_dir, write_outputs

EX = Path(__file__).resolve().parent.parent / "example"
SRC = EX / "数据源20261019.xlsx"


def cells(ws):
    return [[(c.value, c.number_format) for c in row] for row in ws.iter_rows()]


class DefaultOutputDirTest(unittest.TestCase):
    def test_source_stem_plus_timestamp(self):
        src = Path("/tmp") / "会计公司汇总.xlsx"
        now = datetime(2026, 10, 19, 19, 56, 3)
        self.assertEqual(
            default_output_dir([src], now),
            Path("/tmp/会计公司汇总20261019_195603"),
        )


# example/ is git-ignored, so CI may not have it.
@unittest.skipUnless(SRC.exists(), "example files not present")
class ExampleTest(unittest.TestCase):
    def test_matches_expected_outputs(self):
        res = analyze([EX / "数据源20261019.xlsx"])
        self.assertEqual([(g.file_name, g.sheet_name, len(g.rows)) for g in res.groups],
                         [("小明", "A", 11), ("小丁", "B", 9), ("小丁", "D", 11), ("小丁", "E", 10)])
        with tempfile.TemporaryDirectory() as d:
            write_outputs(res.groups, Path(d), Options())
            for name in ("小明.xlsx", "小丁.xlsx"):
                got, exp = load_workbook(Path(d) / name), load_workbook(EX / name)
                self.assertEqual(got.sheetnames, exp.sheetnames)
                for s in exp.sheetnames:
                    self.assertEqual(cells(got[s]), cells(exp[s]), f"{name}/{s}")

    def test_meta_and_stamps(self):
        res = analyze([EX / "数据源20261019.xlsx"])
        opt = Options(include_meta=True, stamp_filename=True, stamp_row=True)
        with tempfile.TemporaryDirectory() as d:
            files = write_outputs(res.groups, Path(d), opt)
            self.assertTrue(all("_20" in f.stem for f in files))
            ws = load_workbook(files[0]).worksheets[0]
            self.assertTrue(ws["A1"].value.startswith("生成时间:"))
            self.assertEqual([ws["A2"].value, ws["B2"].value], ["文件名称", "工作表名称"])
            self.assertEqual([ws["A3"].value, ws["B3"].value], ["小明", "A"])


if __name__ == "__main__":
    unittest.main()
